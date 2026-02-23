import os
import pypdf
import docx
import pptx
import pytesseract
import whisper
from PIL import Image

# IMPORTANT: Set Tesseract path if on Windows and not in PATH
# If you are a student, update this path to where you installed Tesseract
# pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def extract_from_pdf(file_path):
    """
    Extracts text from a PDF file using pypdf.
    """
    text = ""
    try:
        reader = pypdf.PdfReader(file_path)
        for page in reader.pages:
            content = page.extract_text()
            if content:
                text += content + "\n"
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text

def extract_from_docx(file_path):
    """
    Extracts text from a Word document (.docx).
    """
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return text

def extract_from_ppt(file_path):
    """
    Extracts text from a PowerPoint presentation (.pptx).
    """
    text = ""
    try:
        prs = pptx.Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        print(f"Error reading PPTX: {e}")
    return text

def extract_from_txt(file_path):
    """
    Extracts text from a plain text file (.txt).
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        return f"Error reading Text file: {e}"

def extract_from_image(file_path):
    """
    Extracts text from an image using Tesseract OCR.
    Ensure Tesseract is installed on your system!
    """
    try:
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        return f"Error performing OCR: {e}. Is Tesseract installed?"

def extract_from_audio_video(file_path):
    """
    Extracts text from Audio/Video using OpenAI Whisper (Small/Base model).
    This runs on CPU but might be slow for long videos.
    """
    try:
        print("Loading Whisper model... (this might take a moment)")
        # 'tiny' or 'base' is good for CPU. 'small' if you have patience.
        model = whisper.load_model("base") 
        result = model.transcribe(file_path)
        return result["text"]
    except Exception as e:
        return f"Error transcribing audio/video: {e}. Is FFmpeg installed?"
